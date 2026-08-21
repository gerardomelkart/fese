from pathlib import Path
import datetime as dt
import os
import re
import shutil

import numpy as np
import openpyxl
import pandas as pd
import pyreadr
from unidecode import unidecode


BASE_DIR = Path(__file__).resolve().parent
CARPETA_DATOS = BASE_DIR / "datos"
CARPETA_SALIDAS = BASE_DIR / "salidas"
CARPETA_CORRECCIONES = BASE_DIR / "correcciones" / "queretaro"
CARPETA_RESPALDOS = CARPETA_DATOS / "respaldos"

CARPETA_DESTINO = Path(
    r"C:\Users\gerardo.noeller\OneDrive - Secretaría de Seguridad y Protección Ciudadana\Escritorio\FESE"
)

ENTIDAD_OBJETIVO = "QUERETARO"

MESES = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}

MESES_NUMERO = {
    unidecode(nombre).upper(): numero
    for numero, nombre in MESES.items()
}

PERIODOS_ESPERADOS = {
    (2025, 5),
    (2025, 6),
    (2025, 7),
    (2025, 8),
    (2025, 9),
    (2025, 10),
    (2025, 11),
    (2025, 12),
    (2026, 1),
}

COLUMNAS_LONG = [
    "codigo",
    "ao",
    "procedencia",
    "region",
    "estado",
    "centro",
    "tipo",
    "incidente",
    "total",
    "mes",
    "mes_largo",
    "fecha",
]

CODIGOS_IMPROCEDENTES = {
    70101,
    70102,
    70103,
    70104,
    70105,
    70106,
    70107,
    70108,
}


def normalizar_texto(valor):
    return re.sub(
        r"\s+",
        " ",
        unidecode(str(valor)).strip().upper(),
    )


def normalizar_entidad(valor):
    clave = normalizar_texto(valor)

    equivalencias = {
        "QUERETARO": "QUERETARO",
        "QUERETARO DE ARTEAGA": "QUERETARO",
    }

    return equivalencias.get(clave, clave)


def buscar_rds_mas_reciente():
    patron = re.compile(
        r"^fese(\d{4})-(\d{1,2})\.rds$",
        re.IGNORECASE,
    )

    candidatos = []

    for archivo in CARPETA_DATOS.glob("fese*.rds"):
        coincidencia = patron.match(archivo.name)

        if coincidencia:
            periodo = (
                int(coincidencia.group(1)),
                int(coincidencia.group(2)),
            )

            candidatos.append((periodo, archivo))

    if not candidatos:
        raise FileNotFoundError(
            f"No hay históricos RDS en {CARPETA_DATOS}"
        )

    return max(candidatos, key=lambda x: x[0])


def obtener_periodo_archivo(archivo):
    nombre = normalizar_texto(archivo.stem)

    coincidencia = re.search(
        r"FESE[- ]+(ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE)\s+(\d{4})",
        nombre,
    )

    if not coincidencia:
        raise ValueError(
            f"No pude obtener mes y año del archivo: {archivo.name}"
        )

    mes = MESES_NUMERO[coincidencia.group(1)]
    anio = int(coincidencia.group(2))

    return anio, mes


def buscar_archivos_correccion():
    archivos = {}

    for archivo in CARPETA_CORRECCIONES.glob("*.xlsx"):
        if archivo.name.startswith("~$"):
            continue

        periodo = obtener_periodo_archivo(archivo)

        if periodo in archivos:
            raise ValueError(
                f"Hay dos archivos para {periodo}: "
                f"{archivos[periodo].name} y {archivo.name}"
            )

        archivos[periodo] = archivo

    encontrados = set(archivos)

    faltantes = PERIODOS_ESPERADOS - encontrados
    extras = encontrados - PERIODOS_ESPERADOS

    if faltantes:
        texto = ", ".join(
            f"{MESES[mes]} {anio}"
            for anio, mes in sorted(faltantes)
        )

        raise FileNotFoundError(
            f"Faltan archivos corregidos: {texto}"
        )

    if extras:
        texto = ", ".join(
            f"{MESES[mes]} {anio}"
            for anio, mes in sorted(extras)
        )

        raise ValueError(
            f"Hay periodos no esperados en correcciones: {texto}"
        )

    return archivos


def cargar_historico(archivo_rds):
    print(f"Cargando histórico: {archivo_rds.name}")

    historico = pyreadr.read_r(
        str(archivo_rds)
    )[None]

    faltantes = set(COLUMNAS_LONG) - set(historico.columns)

    if faltantes:
        raise ValueError(
            f"Al RDS le faltan columnas: {sorted(faltantes)}"
        )

    historico["codigo"] = pd.to_numeric(
        historico["codigo"],
        errors="raise",
    ).astype(int)

    historico["ao"] = pd.to_numeric(
        historico["ao"],
        errors="raise",
    ).astype(int)

    historico["mes"] = pd.to_numeric(
        historico["mes"],
        errors="raise",
    ).astype(int)

    historico["total"] = pd.to_numeric(
        historico["total"],
        errors="raise",
    )

    return historico


def obtener_catalogos(historico):
    codigos = set(
        historico["codigo"].dropna().astype(int)
    )

    codigo_incidente = (
        historico
        .dropna(subset=["codigo"])
        .groupby("codigo")["incidente"]
        .first()
        .to_dict()
    )

    codigo_tipo = (
        historico
        .dropna(subset=["codigo"])
        .groupby("codigo")["tipo"]
        .first()
        .to_dict()
    )

    clave_estado = historico["estado"].apply(
        normalizar_entidad
    )

    queretaro = historico[
        clave_estado == ENTIDAD_OBJETIVO
    ]

    if queretaro.empty:
        raise ValueError(
            "No se encontraron registros históricos de Querétaro."
        )

    regiones = (
        queretaro["region"]
        .dropna()
        .astype(str)
        .loc[
            lambda x: x.str.strip() != ""
        ]
    )

    if regiones.empty:
        raise ValueError(
            "No fue posible determinar la región de Querétaro."
        )

    region = regiones.mode().iloc[0]

    estados = (
        queretaro["estado"]
        .dropna()
        .astype(str)
    )

    estado = estados.mode().iloc[0]

    return codigos, codigo_incidente, codigo_tipo, region, estado


def valor_fecha(historico, anio, mes):
    fecha = pd.Timestamp(
        year=anio,
        month=mes,
        day=1,
    )

    if pd.api.types.is_datetime64_any_dtype(
        historico["fecha"]
    ):
        return fecha

    muestra = (
        historico["fecha"]
        .dropna()
        .head(1)
    )

    if muestra.empty:
        return fecha.date()

    valor = muestra.iloc[0]

    if isinstance(valor, dt.date) and not isinstance(
        valor,
        dt.datetime,
    ):
        return fecha.date()

    return fecha


def leer_archivo_corregido(
    archivo,
    anio,
    mes,
    historico,
    codigos,
    codigo_incidente,
    codigo_tipo,
    region,
    estado,
):
    print(
        f"Leyendo: {archivo.name}"
    )

    resultados = []

    with pd.ExcelFile(archivo) as xlsx:
        hojas = [
            hoja
            for hoja in xlsx.sheet_names
            if hoja.upper().startswith("C")
        ]

        if not hojas:
            raise ValueError(
                f"{archivo.name} no contiene hojas C."
            )

        datos_hojas = pd.read_excel(
            xlsx,
            sheet_name=hojas,
            usecols="B:D",
            skiprows=18,
        )

        workbook = openpyxl.load_workbook(
            archivo,
            data_only=True,
            read_only=True,
        )

        try:
            resumen_c10 = workbook["RESUMEN"]["C10"].value
            resumen_c11 = workbook["RESUMEN"]["C11"].value

            for hoja in hojas:
                ws = workbook[hoja]

                nombre_centro = ws["D18"].value

                if nombre_centro is None:
                    raise ValueError(
                        f"Centro vacío en {archivo.name}, hoja {hoja}"
                    )

                nombre_centro = str(
                    nombre_centro
                ).strip()

                nombre_entidad = ws["D12"].value

                if nombre_entidad is None:
                    nombre_entidad = resumen_c10

                if nombre_entidad is None:
                    nombre_entidad = resumen_c11

                if nombre_entidad is None:
                    raise ValueError(
                        f"No se encontró entidad en "
                        f"{archivo.name}, hoja {hoja}"
                    )

                if normalizar_entidad(
                    nombre_entidad
                ) != ENTIDAD_OBJETIVO:
                    raise ValueError(
                        f"El archivo {archivo.name}, hoja {hoja}, "
                        f"indica entidad '{nombre_entidad}' "
                        f"y se esperaba Querétaro."
                    )

                datos = datos_hojas[
                    hoja
                ].copy()

                datos.rename(
                    columns={
                        datos.columns[0]: "codigo",
                        datos.columns[1]: "incidente_excel",
                        datos.columns[2]: "total",
                    },
                    inplace=True,
                )

                datos["codigo"] = pd.to_numeric(
                    datos["codigo"],
                    errors="coerce",
                )

                datos["total"] = pd.to_numeric(
                    datos["total"],
                    errors="coerce",
                ).fillna(0)

                datos = datos[
                    datos["codigo"].isin(codigos)
                ].copy()

                datos["codigo"] = (
                    datos["codigo"]
                    .astype(int)
                )

                codigos_sin_catalogo = sorted(
                    set(datos["codigo"])
                    - set(codigo_incidente)
                )

                if codigos_sin_catalogo:
                    raise ValueError(
                        f"Códigos sin incidente histórico en "
                        f"{archivo.name}: {codigos_sin_catalogo}"
                    )

                codigos_sin_tipo = sorted(
                    set(datos["codigo"])
                    - set(codigo_tipo)
                )

                if codigos_sin_tipo:
                    raise ValueError(
                        f"Códigos sin tipo histórico en "
                        f"{archivo.name}: {codigos_sin_tipo}"
                    )

                datos["ao"] = anio
                datos["procedencia"] = np.where(
                    datos["codigo"].isin(
                        CODIGOS_IMPROCEDENTES
                    ),
                    "Improcedentes",
                    "Procedentes",
                )

                datos["region"] = region
                datos["estado"] = estado
                datos["centro"] = nombre_centro

                datos["tipo"] = datos[
                    "codigo"
                ].map(codigo_tipo)

                datos["incidente"] = datos[
                    "codigo"
                ].map(codigo_incidente)

                datos["mes"] = mes
                datos["mes_largo"] = MESES[mes]
                datos["fecha"] = valor_fecha(
                    historico,
                    anio,
                    mes,
                )

                datos = datos[
                    COLUMNAS_LONG
                ]

                resultados.append(datos)

        finally:
            workbook.close()

    if not resultados:
        raise ValueError(
            f"No se obtuvieron datos de {archivo.name}"
        )

    salida = pd.concat(
        resultados,
        ignore_index=True,
    )

    return salida


def construir_correcciones(
    historico,
    archivos,
):
    (
        codigos,
        codigo_incidente,
        codigo_tipo,
        region,
        estado,
    ) = obtener_catalogos(historico)

    correcciones = {}

    for (anio, mes), archivo in sorted(
        archivos.items()
    ):
        correcciones[
            (anio, mes)
        ] = leer_archivo_corregido(
            archivo=archivo,
            anio=anio,
            mes=mes,
            historico=historico,
            codigos=codigos,
            codigo_incidente=codigo_incidente,
            codigo_tipo=codigo_tipo,
            region=region,
            estado=estado,
        )

    return correcciones


def crear_auditoria(
    historico,
    correcciones,
):
    clave_estado = historico[
        "estado"
    ].apply(normalizar_entidad)

    filas = []

    for (anio, mes), nuevos in sorted(
        correcciones.items()
    ):
        mascara = (
            (clave_estado == ENTIDAD_OBJETIVO)
            & (historico["ao"] == anio)
            & (historico["mes"] == mes)
        )

        anteriores = historico[
            mascara
        ]

        if anteriores.empty:
            raise ValueError(
                f"No existen registros actuales de Querétaro "
                f"para {MESES[mes]} {anio}."
            )

        total_anterior = float(
            anteriores["total"].sum()
        )

        total_corregido = float(
            nuevos["total"].sum()
        )

        filas.append(
            {
                "periodo": f"{MESES[mes]} {anio}",
                "anio": anio,
                "mes": mes,
                "registros_anteriores": len(
                    anteriores
                ),
                "registros_corregidos": len(
                    nuevos
                ),
                "total_anterior": total_anterior,
                "total_corregido": total_corregido,
                "diferencia": (
                    total_corregido
                    - total_anterior
                ),
            }
        )

    return pd.DataFrame(filas)


def mostrar_auditoria(auditoria):
    print()
    print(
        "===== COMPARATIVO ANTES DE APLICAR ====="
    )
    print()

    for fila in auditoria.itertuples():
        print(
            f"{fila.periodo:<18} "
            f"Anterior: {fila.total_anterior:>12,.0f}   "
            f"Corregido: {fila.total_corregido:>12,.0f}   "
            f"Diferencia: {fila.diferencia:>+12,.0f}"
        )

    print()


def aplicar_correcciones(
    historico,
    correcciones,
):
    clave_estado = historico[
        "estado"
    ].apply(normalizar_entidad)

    eliminar = pd.Series(
        False,
        index=historico.index,
    )

    for anio, mes in correcciones:
        eliminar |= (
            (clave_estado == ENTIDAD_OBJETIVO)
            & (historico["ao"] == anio)
            & (historico["mes"] == mes)
        )

    print(
        f"Registros históricos de Querétaro "
        f"a sustituir: {int(eliminar.sum()):,}"
    )

    conservados = historico[
        ~eliminar
    ].copy()

    nuevos = pd.concat(
        [
            correcciones[periodo]
            for periodo in sorted(
                correcciones
            )
        ],
        ignore_index=True,
    )

    corregido = pd.concat(
        [
            conservados[COLUMNAS_LONG],
            nuevos[COLUMNAS_LONG],
        ],
        ignore_index=True,
    )

    return corregido


def validar_resultado(
    corregido,
    correcciones,
):
    clave_estado = corregido[
        "estado"
    ].apply(normalizar_entidad)

    for (anio, mes), esperado in correcciones.items():
        actual = corregido[
            (clave_estado == ENTIDAD_OBJETIVO)
            & (corregido["ao"] == anio)
            & (corregido["mes"] == mes)
        ]

        if len(actual) != len(esperado):
            raise RuntimeError(
                f"Validación fallida en {MESES[mes]} {anio}: "
                f"se esperaban {len(esperado)} registros "
                f"y quedaron {len(actual)}."
            )

        total_actual = float(
            actual["total"].sum()
        )

        total_esperado = float(
            esperado["total"].sum()
        )

        if abs(
            total_actual - total_esperado
        ) > 0.001:
            raise RuntimeError(
                f"Validación fallida en {MESES[mes]} {anio}: "
                f"{total_actual} != {total_esperado}"
            )


def respaldar_archivo(
    archivo,
    marca,
):
    CARPETA_RESPALDOS.mkdir(
        parents=True,
        exist_ok=True,
    )

    destino = (
        CARPETA_RESPALDOS
        / f"{archivo.stem}_ANTES_CORRECCION_QUERETARO_{marca}{archivo.suffix}"
    )

    print(
        f"Generando respaldo: {destino.name}"
    )

    shutil.copy2(
        archivo,
        destino,
    )

    return destino


def escribir_rds_seguro(
    dataframe,
    archivo_rds,
):
    temporal = archivo_rds.with_name(
        f"~TEMP_{archivo_rds.name}"
    )

    if temporal.exists():
        temporal.unlink()

    print(
        "Escribiendo RDS corregido..."
    )

    pyreadr.write_rds(
        str(temporal),
        dataframe,
    )

    os.replace(
        temporal,
        archivo_rds,
    )

    print(
        f"RDS actualizado: {archivo_rds}"
    )


def generar_reporte_anual(
    historico,
    anio_corte,
    mes_corte,
):
    print(
        "Generando reporte anual corregido..."
    )

    datos = historico[
        COLUMNAS_LONG
    ].copy()

    wide = datos.pivot(
        index=[
            "ao",
            "region",
            "estado",
            "centro",
            "tipo",
            "codigo",
            "incidente",
            "procedencia",
        ],
        columns="mes_largo",
        values="total",
    )

    wide = wide.reset_index()

    wide.rename(
        columns={
            "codigo": "Código",
            "ao": "Año",
            "procedencia": "Procedencia",
            "region": "Region",
            "estado": "Estado",
            "centro": "Centro",
            "tipo": "Tipo",
            "incidente": "Incidente",
        },
        inplace=True,
    )

    columnas = [
        "Código",
        "Año",
        "Procedencia",
        "Region",
        "Estado",
        "Centro",
        "Tipo",
        "Incidente",
        "Enero",
        "Febrero",
        "Marzo",
        "Abril",
        "Mayo",
        "Junio",
        "Julio",
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
        "Diciembre",
    ]

    for mes in MESES.values():
        if mes not in wide.columns:
            wide[mes] = 0
        else:
            wide[mes] = (
                wide[mes]
                .fillna(0)
            )

    wide = wide[
        columnas
    ]

    wide["Estado"] = (
        wide["Estado"]
        .astype(str)
        .str.upper()
    )

    wide = wide.sort_values(
        by=[
            "Año",
            "Código",
        ],
        ascending=[
            True,
            True,
        ],
    )

    CARPETA_SALIDAS.mkdir(
        parents=True,
        exist_ok=True,
    )

    archivo = (
        CARPETA_SALIDAS
        / f"Rep_anual{anio_corte}-{mes_corte:02d}.xlsx"
    )

    wide.to_excel(
        archivo,
        index=False,
    )

    print(
        f"Reporte anual actualizado: {archivo}"
    )

    return archivo


def generar_auditoria_excel(
    auditoria,
    marca,
):
    CARPETA_SALIDAS.mkdir(
        parents=True,
        exist_ok=True,
    )

    archivo = (
        CARPETA_SALIDAS
        / f"auditoria_correccion_queretaro_{marca}.xlsx"
    )

    with pd.ExcelWriter(
        archivo,
        engine="openpyxl",
    ) as writer:
        auditoria.to_excel(
            writer,
            index=False,
            sheet_name="Comparativo",
        )

        hoja = writer.book[
            "Comparativo"
        ]

        hoja.freeze_panes = "A2"

        anchos = {
            "A": 20,
            "B": 10,
            "C": 10,
            "D": 22,
            "E": 22,
            "F": 20,
            "G": 20,
            "H": 20,
        }

        for columna, ancho in anchos.items():
            hoja.column_dimensions[
                columna
            ].width = ancho

    print(
        f"Auditoría generada: {archivo}"
    )

    return archivo


def copiar_reemplazo_seguro(
    origen,
    destino,
):
    destino.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporal = destino.with_name(
        f"~TEMP_{destino.name}"
    )

    if temporal.exists():
        temporal.unlink()

    shutil.copy2(
        origen,
        temporal,
    )

    os.replace(
        temporal,
        destino,
    )


def actualizar_onedrive(
    archivo_rds,
    archivo_excel,
):
    print(
        "Actualizando copias de OneDrive..."
    )

    copiar_reemplazo_seguro(
        archivo_rds,
        CARPETA_DESTINO
        / archivo_rds.name,
    )

    copiar_reemplazo_seguro(
        archivo_excel,
        CARPETA_DESTINO
        / archivo_excel.name,
    )

    print(
        "Copias de OneDrive actualizadas."
    )


def limpiar_cache(
    archivo_rds,
):
    localappdata = os.getenv(
        "LOCALAPPDATA"
    )

    if not localappdata:
        return

    carpeta_cache = (
        Path(localappdata)
        / "FESE"
        / "cache"
    )

    if not carpeta_cache.is_dir():
        return

    stem = archivo_rds.stem

    for archivo in carpeta_cache.glob(
        f"{stem}*"
    ):
        if archivo.is_file():
            archivo.unlink()

            print(
                f"Cache eliminado: {archivo.name}"
            )


def main():
    print(
        "===== CORRECCIÓN HISTÓRICA FESE - QUERÉTARO ====="
    )

    CARPETA_CORRECCIONES.mkdir(
        parents=True,
        exist_ok=True,
    )

    (
        periodo_rds,
        archivo_rds,
    ) = buscar_rds_mas_reciente()

    print(
        f"Histórico que se modificará: {archivo_rds}"
    )

    archivos = buscar_archivos_correccion()

    print(
        f"Archivos corregidos encontrados: {len(archivos)}"
    )

    historico = cargar_historico(
        archivo_rds
    )

    correcciones = construir_correcciones(
        historico,
        archivos,
    )

    auditoria = crear_auditoria(
        historico,
        correcciones,
    )

    mostrar_auditoria(
        auditoria
    )

    print(
        "Se sustituirán únicamente los registros de Querétaro "
        "de mayo a diciembre de 2025 y enero de 2026."
    )

    confirmar = input(
        "\nEscribe APLICAR para continuar: "
    ).strip().upper()

    if confirmar != "APLICAR":
        print(
            "Operación cancelada. No se modificó ningún archivo."
        )

        return

    marca = dt.datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    respaldo_rds = respaldar_archivo(
        archivo_rds,
        marca,
    )

    corregido = aplicar_correcciones(
        historico,
        correcciones,
    )

    validar_resultado(
        corregido,
        correcciones,
    )

    escribir_rds_seguro(
        corregido,
        archivo_rds,
    )

    limpiar_cache(
        archivo_rds
    )

    archivo_excel = generar_reporte_anual(
        corregido,
        periodo_rds[0],
        periodo_rds[1],
    )

    archivo_auditoria = generar_auditoria_excel(
        auditoria,
        marca,
    )

    actualizar_onedrive(
        archivo_rds,
        archivo_excel,
    )

    print()
    print(
        "===== CORRECCIÓN TERMINADA CORRECTAMENTE ====="
    )
    print(
        f"RDS actualizado: {archivo_rds}"
    )
    print(
        f"Respaldo original: {respaldo_rds}"
    )
    print(
        f"Reporte anual: {archivo_excel}"
    )
    print(
        f"Auditoría: {archivo_auditoria}"
    )


if __name__ == "__main__":
    main()